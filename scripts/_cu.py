# -*- coding: utf-8 -*-
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from docx import Document
D = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\STF_Gomez_Javier_E1_v1.docx"
doc = Document(D)
ps = doc.paragraphs
print("parrafos:", len(ps), "tablas:", len(doc.tables))
ini = next(i for i,p in enumerate(ps) if "10. PLAN DE DESARROLLO" in p.text)
print("inicio PDT:", ini)
# titulos 10.x
for i in range(ini, len(ps)):
    t = ps[i].text.strip()
    if re.match(r'^10\.\d+(\.\d+)*\s', t) and len(t) < 90:
        print(f"  [{i}] {t}")
