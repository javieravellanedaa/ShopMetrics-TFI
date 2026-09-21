# -*- coding: utf-8 -*-
import openpyxl,sys,io,re
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding="utf-8",errors="replace")
p=r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wb=openpyxl.load_workbook(p,data_only=True)
pat=re.compile(r"sensor|contador|kit|equipam|hardware|gateway|dispositiv", re.I)
for sn in wb.sheetnames:
    ws=wb[sn]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,str) and pat.search(c.value):
                vec=[]
                for cc in range(c.column, min(c.column+5, ws.max_column+1)):
                    v=ws.cell(c.row,cc).value
                    if v is not None: vec.append(f"{ws.cell(c.row,cc).coordinate}={repr(v)[:46]}")
                print(f"{sn}: "+" | ".join(vec))
