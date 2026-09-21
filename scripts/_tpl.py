# -*- coding: utf-8 -*-
import openpyxl,sys,io
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding="utf-8",errors="replace")
for lab,p in [("TEMPLATE V1", r"D:\Javier\UAI\UAI - SAP\STF\STF\Ejemplos de catedra\Presupuesto financiero EJEMPLO V1.xlsx"),
              ("BK preriesgos", r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Backups\Presupuesto financiero ShopMetrics V1.BACKUP-preriesgos.xlsx")]:
    wb=openpyxl.load_workbook(p); ws=wb["Anexo capacidad operativa"]
    print("="*70); print(lab, " dims:", ws.max_row, ws.max_column)
    print("merges:", sorted(str(m) for m in ws.merged_cells.ranges)[:60])
    print("widths:", {k:round(v.width,1) for k,v in sorted(ws.column_dimensions.items()) if v.width})
    n=0
    for r in range(1,120):
        for c in range(1,30):
            cf=ws.cell(r,c)
            if cf.value is None: continue
            n+=1
            if n<=55: print(f"   {cf.coordinate} = {repr(cf.value)[:75]}")
    print("  celdas con contenido:", n)
