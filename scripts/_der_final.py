# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from PIL import Image

D   = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\STF_Gomez_Javier_E1_v1.docx"
IMG = r"D:\Javier\UAI\UAI - SAP\shopMetrics\shopMetrics\enterprise_architect\DER_ShopMetrics_completo.png"
doc = Document(D)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
def blocks(parent):
    for ch in parent.element.body.iterchildren():
        if ch.tag == W+"p":  yield Paragraph(ch, parent)
        elif ch.tag == W+"tbl": yield Table(ch, parent)
seq = list(blocks(doc))
ini = [i for i,b in enumerate(seq) if isinstance(b,Paragraph) and b.text.strip().startswith("10.5.8")][-1]
p_img, p_cap = seq[ini+2], seq[ini+3]

for r in list(p_img.runs): r._r.getparent().remove(r._r)
w, h = Image.open(IMG).size
p_img.alignment = AL.CENTER
p_img.paragraph_format.keep_with_next = True
p_img.paragraph_format.space_after = Pt(4)
p_img.add_run().add_picture(IMG, width=Cm(16.26))
print(f"DER completo: 16.26 x {16.26*h/w:.2f} cm")

nuevo = ("Figura. Diagrama Entidad-Relación completo de ShopMetrics: las 30 entidades con todas sus "
         "columnas y tipos de dato, las claves primarias y foráneas, y las 43 relaciones en notación de "
         "pata de gallo. Por la cantidad de entidades, la lectura en tamaño real se hace sobre el archivo "
         "«DER_ShopMetrics_completo_A3.pdf» (A3 apaisado) o sobre los ocho diagramas por área que siguen. "
         "Fuente: Enterprise Architect.")
for r in list(p_cap.runs)[1:]: r._r.getparent().remove(r._r)
if p_cap.runs: p_cap.runs[0].text = nuevo
else:
    r = p_cap.add_run(nuevo); r.italic = True; r.font.size = Pt(9)
doc.save(D)
print("Word actualizado")
