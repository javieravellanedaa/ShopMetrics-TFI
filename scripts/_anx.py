# -*- coding: utf-8 -*-
import openpyxl
p = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wf = openpyxl.load_workbook(p, data_only=False)
wv = openpyxl.load_workbook(p, data_only=True)
n = [s for s in wf.sheetnames if "apacidad" in s][0]
print("HOJA:", n)
f, v = wf[n], wv[n]
print("merges:", sorted(str(m) for m in f.merged_cells.ranges))
print("dims:", f.max_row, f.max_column)
for r in range(1, min(f.max_row,70)+1):
    fila=[]
    for c in range(1, min(f.max_column,9)+1):
        cf=f.cell(r,c); cv=v.cell(r,c)
        if cf.value is None and cv.value is None: continue
        s=f"{cf.coordinate}="
        if isinstance(cf.value,str) and cf.value.startswith("="):
            s+=f"{cf.value} -> {cv.value}"
        else:
            s+=repr(cf.value)
        fila.append(s)
    if fila: print(f"[{r:>2}]", " | ".join(fila))
