# -*- coding: utf-8 -*-
"""Reemplaza la figura del DER y agrega los 8 diagramas por area."""
import sys, io, os, copy, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from PIL import Image

D   = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\STF_Gomez_Javier_E1_v1.docx"
EA  = r"D:\Javier\UAI\UAI - SAP\shopMetrics\shopMetrics\enterprise_architect"
SUB = os.path.join(EA, "export", "der")
MASTER = os.path.join(EA, "DER_ShopMetrics.png")

ANCHO_MAX, ALTO_MAX = 16.26, 18.0
def medida(p):
    w, h = Image.open(p).size
    a = min(ANCHO_MAX, ALTO_MAX * w / h)
    return a, a * h / w

doc = Document(D)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
def blocks(parent):
    for ch in parent.element.body.iterchildren():
        if ch.tag == W+"p":  yield Paragraph(ch, parent)
        elif ch.tag == W+"tbl": yield Table(ch, parent)
seq = list(blocks(doc))
ini = [i for i,b in enumerate(seq) if isinstance(b,Paragraph) and b.text.strip().startswith("10.5.8")][-1]

p_intro = seq[ini+1]     # texto introductorio
p_img   = seq[ini+2]     # parrafo con la imagen vieja del DER
p_cap   = seq[ini+3]     # leyenda

# ---------- 1) texto introductorio y leyenda ----------
for p in (p_intro, p_cap):
    for r in p.runs:
        r.text = r.text.replace("25 entidades", "30 entidades")
txt = "".join(r.text for r in p_intro.runs)
if "áreas" not in txt:
    p_intro.add_run(" Por su tamaño, el modelo se presenta primero en una vista general y luego "
                    "desagregado en ocho diagramas por área temática, donde cada relación se lee con "
                    "notación de pata de gallo (Information Engineering).")
print("intro y leyenda actualizados")

# ---------- 2) reemplazar la imagen del DER ----------
for r in list(p_img.runs):
    r._r.getparent().remove(r._r)
a, alto = medida(MASTER)
p_img.alignment = AL.CENTER
p_img.paragraph_format.keep_with_next = True
p_img.paragraph_format.space_after = Pt(4)
p_img.add_run().add_picture(MASTER, width=Cm(a))
print(f"DER maestro insertado: {a:.2f} x {alto:.2f} cm")

# ---------- 3) los 8 diagramas por area, despues de la leyenda ----------
FIGS = [
 ("DER_1_estructura.png",     "Estructura del centro: centros, zonas, locatarios, locales y consentimientos."),
 ("DER_2_seguridad.png",      "Seguridad y acceso: usuarios, roles, permisos y acceso multicentro."),
 ("DER_3_integraciones.png",  "Integraciones e ingesta: conectores POS, dispositivos IoT y los datos que producen."),
 ("DER_4_reglas.png",         "Reglas de umbral, destinatarios y generación de alertas por regla o por modelo de ML."),
 ("DER_5_resolucion.png",     "Resolución de alertas: asignación al operario, causa y acción tomada."),
 ("DER_6_ml.png",             "Modelos de ML: predicción de vacancia, acciones de retención y recomendaciones."),
 ("DER_7_auditoria.png",      "Auditoría y reportes, ambos acotados al centro sobre el que se opera."),
 ("DER_8_notificaciones.png", "Notificaciones, preferencias del usuario y parámetros del sistema."),
]
anc = p_cap._p
n = 8
for arch, leyenda in FIGS:
    ruta = os.path.join(SUB, arch)
    a, alto = medida(ruta)
    pi = doc.add_paragraph(); anc.addnext(pi._p); anc = pi._p
    pi.alignment = AL.CENTER
    pi.paragraph_format.keep_with_next = True
    pi.paragraph_format.space_after = Pt(4)
    pi.add_run().add_picture(ruta, width=Cm(a))
    pc = doc.add_paragraph(); anc.addnext(pc._p); anc = pc._p
    pc.alignment = AL.CENTER
    r = pc.add_run(f"Figura. {leyenda} Fuente: Enterprise Architect.")
    r.italic = True; r.font.size = Pt(9)
    print(f"  + {arch}  {a:.1f} x {alto:.1f} cm")

doc.save(D)
print("figuras listas")
