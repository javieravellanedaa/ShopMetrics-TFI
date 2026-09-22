# -*- coding: utf-8 -*-
import openpyxl,sys,io
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding="utf-8",errors="replace")
p=r"D:\Javier\UAI\UAI - SAP\STF\STF\Ejemplos de catedra\Ejemplo PresupuestoFinanciero.xlsx"
wf=openpyxl.load_workbook(p,data_only=False); wv=openpyxl.load_workbook(p,data_only=True)
n=[s for s in wf.sheetnames if "apacidad" in s][0]
f,v=wf[n],wv[n]
print("EJEMPLO CATEDRA -", n, "dims:", f.max_row, f.max_column)
for r in range(32, min(f.max_row,80)+1):
    fila=[]
    for c in range(1, min(f.max_column,14)+1):
        cf=f.cell(r,c); cv=v.cell(r,c)
        if cf.value is None and cv.value is None: continue
        s=f"{cf.coordinate}="
        if isinstance(cf.value,str) and cf.value.startswith("="): s+=f"{cf.value}->{cv.value}"
        else: s+=repr(cf.value)[:110]
        fila.append(s)
    if fila: print(f"[{r:>2}]"," | ".join(fila))
