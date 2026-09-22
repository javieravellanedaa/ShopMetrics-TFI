# -*- coding: utf-8 -*-
"""Carga los datos de ShopMetrics en el Anexo SIN tocar el formato del template.
Solo se escribe .value: estilos, merges, anchos, altos y colores quedan intactos."""
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

P = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wb = openpyxl.load_workbook(P); ws = wb["Anexo capacidad operativa"]

def S(cell, val):
    ws[cell].value = val

# ------------- encabezados de los bloques -------------
S("A3", "Instalación y alta de los planes ShopMetrics")
S("H3", "Soporte mensual por plan")
S("P3", 6)                                   # costo del tecnico por hora, USD
S("S3", "Coste del kit de instalación")
S("Y3", "Coste del soporte mensual")
S("T11", "Alta remota: conexión con el punto de venta")
S("T12", "Visita al local: sensor de vidriera + contador de puerta")
S("T13", "Instalación en hasta cinco locales de una cadena")

# ------------- bloques de tareas -------------
# (fila_titulo, titulo_instalacion, tareas_instalacion, titulo_soporte, tareas_soporte)
BLOQUES = [
 (18, "Alta e instalación Plan Básico",
  [("Contacto y coordinación con el comercio",             0.08, 1),
   ("Alta del comercio en la plataforma",                  0.08, 1),
   ("Conexión con el sistema de punto de venta",           0.17, 1),
   ("Verificación de la calidad de los datos",             0.12, 1),
   ("Capacitación remota del comerciante",                 0.05, 1)],
  "Soporte mensual Plan Básico",
  [("Monitoreo automático de la ingesta de datos",         0.05, 1),
   ("Atención de consultas por canal remoto",              0.07, 1),
   ("Revisión y envío del reporte mensual",                0.03, 1)]),
 (37, "Alta e instalación Plan Vidriera",
  [("Coordinación de la visita y preparación del kit",     0.25, 1),
   ("Traslado al local",                                   0.50, 1),
   ("Montaje del sensor de vidriera",                      0.75, 1),
   ("Instalación del contador de puerta",                  0.50, 1),
   ("Conexión con el sistema de punto de venta",           0.25, 1),
   ("Prueba del embudo y calibración",                     0.50, 1),
   ("Capacitación del comerciante y cierre",               0.25, 1)],
  "Soporte mensual Plan Vidriera",
  [("Monitoreo automático de la ingesta de datos",         0.05, 1),
   ("Control del estado del sensor de vidriera",           0.07, 1),
   ("Atención de consultas por canal remoto",              0.08, 1),
   ("Revisión y envío del reporte mensual",                0.05, 1)]),
 (56, "Alta e instalación Plan Cadena",
  [("Relevamiento previo de la cadena",                    1.00, 1),
   ("Alta de la cadena y de sus sucursales",               0.50, 1),
   ("Traslado entre locales",                              0.50, 5),
   ("Montaje del sensor de vidriera",                      0.75, 5),
   ("Instalación del contador de puerta",                  0.40, 5),
   ("Conexión del punto de venta por local",               0.20, 5),
   ("Consolidación y prueba comparativa entre sucursales", 0.75, 1),
   ("Capacitación del responsable de la cadena",           0.50, 1)],
  "Soporte mensual Plan Cadena",
  [("Monitoreo automático de la ingesta (5 locales)",      0.04, 5),
   ("Control del estado de los sensores",                  0.03, 5),
   ("Atención de consultas del responsable de cadena",     0.10, 1),
   ("Reporte comparativo entre sucursales",                0.15, 1)]),
]

for r0, tit_i, tar_i, tit_s, tar_s in BLOQUES:
    S(f"A{r0}", tit_i)
    S(f"H{r0}", tit_s)
    for i in range(12):                       # 12 renglones del template
        r = r0 + 2 + i
        d, bh, ct = tar_i[i] if i < len(tar_i) else (None, None, None)
        S(f"B{r}", d); S(f"C{r}", bh); S(f"D{r}", ct)
        d, bh, ct = tar_s[i] if i < len(tar_s) else (None, None, None)
        S(f"I{r}", d); S(f"J{r}", bh); S(f"K{r}", ct)

# ------------- comentarios de los bloques -------------
S("A32", "Comentarios: el alta del Plan Básico es remota, no requiere visita al local.")
S("A51", "Comentarios: incluye el traslado y el montaje del equipamiento en vidriera.")
S("A70", "Comentarios: cubre hasta cinco locales de una misma cadena.")
S("H32", "Comentarios: el soporte del Plan Básico es enteramente remoto.")
S("H51", "Comentarios: suma el control del estado del sensor de vidriera.")
S("H70", "Comentarios: el monitoreo se repite por cada local de la cadena.")
S("A76", "Detalle: el soporte por plan es mensual y se sostiene mientras el comercio permanece abonado.")

# ------------- kits de equipamiento (columnas S/U/V) -------------
KIT = {18: [("Instalación (mano de obra)", "=$E$32", "=$P$3")],
       37: [("Sensor de vidriera",      1, 14),
            ("Contador de puerta",      1,  9),
            ("Gateway de conexión",     1,  7),
            ("Kit de montaje",          1,  4),
            ("Instalación (mano de obra)", "=$E$51", "=$P$3")],
       56: [("Sensor de vidriera",      5, 14),
            ("Contador de puerta",      5,  9),
            ("Gateway de conexión",     5,  7),
            ("Kit de montaje",          5,  4),
            ("Instalación (mano de obra)", "=$E$70", "=$P$3")]}
TIT_KIT = {18: "Kit Plan Básico", 37: "Kit Plan Vidriera", 56: "Kit Plan Cadena"}
FIN     = {18: 27, 37: 46, 56: 64}            # ultima fila util de cada bloque de kit
for r0, comps in KIT.items():
    S(f"S{r0}", TIT_KIT[r0])
    S(f"Y{r0}", TIT_KIT[r0].replace("Kit", "Soporte mensual"))
    r = r0 + 2
    for nom, cant, pu in comps:
        S(f"S{r}", nom); S(f"U{r}", cant); S(f"V{r}", pu)
        ws[f"W{r}"].value = f"=V{r}*U{r}"
        r += 1
    for rr in range(r, FIN[r0] + 1):          # limpiar los renglones sobrantes
        S(f"S{rr}", None); S(f"U{rr}", None); S(f"V{rr}", None); S(f"W{rr}", None)

# soporte mensual valorizado (columnas Y/AA/AB)
for r0, ref in ((18, "$L$32"), (37, "$L$51"), (56, "$L$70")):
    S(f"Y{r0+2}", "Servicio técnico")
    ws[f"AA{r0+2}"].value = f"={ref}"
    ws[f"AB{r0+2}"].value = "=$P$3"

# ------------- tablas anuales: filas de abonos -------------
MES = ["C","E","G","I","K","M","O","Q","S","U","W","Y"]      # columnas del anexo
PV  = ["E","G","I","K","M","O","Q","S","U","W","Y","AA"]     # columnas de Proy. ventas
# (fila_anexo, fila_Proy.ventas de abonos)
ABONOS = [(83, 22), (84, 23), (85, 24),        # 2026
          (94, 85), (95, 86), (96, 87),        # 2027
          (105, 147), (106, 148), (107, 149)]  # 2028
for ra, rp in ABONOS:
    for ca, cp in zip(MES, PV):
        ws[f"{ca}{ra}"].value = f"='Proy. ventas'!{cp}{rp}"

# ------------- conclusion -------------
S("A111",
  "El anexo dimensiona la capacidad operativa desde la demanda. Cada plan tiene un consumo horario propio, obtenido de la suma de "
  "las tareas que requiere su alta, y un consumo mensual de soporte que se sostiene mientras el comercio permanece abonado: "
  "0,50 h el Plan Básico, 3,00 h el Plan Vidriera y 12,00 h el Plan Cadena para el alta; 0,15 h, 0,25 h y 0,60 h al mes para el soporte.\n\n"
  "Multiplicando esas bases horarias por las unidades de la proyección de ventas se obtienen las horas hombre requeridas mes a mes. "
  "El total asciende a 440,75 horas en 2026, 1.113,75 en 2027 y 1.982,10 en 2028, con picos mensuales de 87,50, 123,35 y 189,65 horas "
  "respectivamente. Contra las 150 horas laborables mensuales de un empleado, la dotación técnica necesaria es de un técnico en 2026 "
  "y 2027, y de dos a partir de marzo de 2028. El mes en que ese cociente supera la unidad es el que marca la incorporación en la "
  "estructura de costos de recursos humanos: la contratación no se decide por criterio, se desprende del volumen proyectado.\n\n"
  "El costo de mano de obra por alta resulta de USD 3 en el Plan Básico, USD 18 en el Plan Vidriera y USD 72 en el Plan Cadena, contra "
  "precios de lista de USD 25, USD 60 y USD 250. Sumado el kit de equipamiento, el alta cubre el esfuerzo técnico comprometido en los "
  "tres planes.")

wb.save(P); print("OK datos cargados")
