# -*- coding: utf-8 -*-
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from docx import Document
D = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\STF_Gomez_Javier_E1_v1.docx"
doc = Document(D)
ps = doc.paragraphs
print("total parrafos:", len(ps), " tablas:", len(doc.tables))
# ubicar los headings
for i, p in enumerate(ps):
    st = (p.style.name or "").lower()
    t = p.text.strip()
    if st.startswith("heading") or st.startswith("título") or st.startswith("titulo"):
        if t: print(f"[{i:>5}] {st:<12} {t[:90]}")
