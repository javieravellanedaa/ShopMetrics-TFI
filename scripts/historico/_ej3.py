# -*- coding: utf-8 -*-
import openpyxl,sys,io
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding="utf-8",errors="replace")
p=r"D:\Javier\UAI\UAI - SAP\STF\STF\Ejemplos de catedra\Ejemplo PresupuestoFinanciero.xlsx"
wf=openpyxl.load_workbook(p,data_only=False); wv=openpyxl.load_workbook(p,data_only=True)
f,v=wf["Anexo capacidad operativa"],wv["Anexo capacidad operativa"]
for r in range(77, 130):
    fila=[]
    for c in range(1, 30):
        cf=f.cell(r,c); cv=v.cell(r,c)
        if cf.value is None and cv.value is None: continue
        s=f"{cf.coordinate}="
        if isinstance(cf.value,str) and cf.value.startswith("="): s+=f"{cf.value}->{cv.value}"
        else: s+=repr(cf.value)[:70]
        fila.append(s)
    if fila: print(f"[{r:>3}]"," | ".join(fila))
