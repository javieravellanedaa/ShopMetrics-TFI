# -*- coding: utf-8 -*-
"""Acorta las descripciones para que entren en los anchos del template (B=31,6 / I=22,4).
No se toca ningun atributo de formato: solo .value."""
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
P = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wb = openpyxl.load_workbook(P); ws = wb["Anexo capacidad operativa"]

INST = {18: ["Contacto y coordinación", "Alta en la plataforma", "Conexión con el punto de venta",
             "Verificación de los datos", "Capacitación remota"],
        37: ["Coordinación y armado del kit", "Traslado al local", "Montaje del sensor de vidriera",
             "Instalación contador de puerta", "Conexión con el punto de venta",
             "Prueba del embudo", "Capacitación y cierre"],
        56: ["Relevamiento de la cadena", "Alta de cadena y sucursales", "Traslado entre locales",
             "Montaje sensor de vidriera", "Instalación contador de puerta",
             "Conexión del punto de venta", "Prueba comparativa", "Capacitación del responsable"]}
SOP  = {18: ["Monitoreo de ingesta", "Atención de consultas", "Reporte mensual"],
        37: ["Monitoreo de ingesta", "Control del sensor", "Atención de consultas", "Reporte mensual"],
        56: ["Monitoreo por local", "Control de sensores", "Consultas del titular", "Reporte comparativo"]}

for r0 in (18, 37, 56):
    for i, t in enumerate(INST[r0]):
        ws[f"B{r0+2+i}"].value = t
    for i, t in enumerate(SOP[r0]):
        ws[f"I{r0+2+i}"].value = t

# comentarios: caben en el bloque A:C fusionado, se acortan igual
ws["A32"].value = "Comentarios: el alta del Plan Básico es remota, no requiere visita."
ws["A51"].value = "Comentarios: incluye traslado y montaje del equipamiento en vidriera."
ws["A70"].value = "Comentarios: cubre hasta cinco locales de una misma cadena."
ws["H32"].value = "Comentarios: soporte enteramente remoto."
ws["H51"].value = "Comentarios: suma el control del sensor de vidriera."
ws["H70"].value = "Comentarios: el monitoreo se repite por local."

wb.save(P)
print("largos B:", [len(t) for r in INST.values() for t in r])
print("largos I:", [len(t) for r in SOP.values() for t in r])
