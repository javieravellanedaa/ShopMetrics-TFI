# -*- coding: utf-8 -*-
"""Reconstruye 'Anexo capacidad operativa' replicando la estructura del ejemplo de cátedra."""
import openpyxl, sys, io
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

P = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wb = openpyxl.load_workbook(P)
ws = wb["Anexo capacidad operativa"]

# ---------- paleta ----------
VERDE  = PatternFill("solid", fgColor="A8D08D")   # titulo de bloque
VERDE2 = PatternFill("solid", fgColor="C5E0B3")   # subtitulo
HDR    = PatternFill("solid", fgColor="E2EFD9")   # encabezado de tabla
AMAR   = PatternFill("solid", fgColor="FFF2CC")   # dato de entrada
AZUL   = PatternFill("solid", fgColor="D9E2F3")   # dato calculado / cuerpo
GRIS   = PatternFill("solid", fgColor="F2F2F2")
TOTFIL = PatternFill("solid", fgColor="FFD966")   # fila de totales

thin = Side(style="thin", color="BFBFBF")
med  = Side(style="medium", color="808080")
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)
BOXB = Border(left=thin, right=thin, top=thin, bottom=med)

NUM2 = "#,##0.00"
NUM0 = "#,##0"
USD  = '_ [$$-2C0A]\\ * #,##0.00_ ;_ [$$-2C0A]\\ * \\-#,##0.00_ ;_ [$$-2C0A]\\ * "-"??_ ;_ @_ '
CTR  = Alignment(horizontal="center", vertical="center", wrap_text=True)
LFT  = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RGT  = Alignment(horizontal="right",  vertical="center")

def B(cell, val, fmt=None, fill=None, bold=False, align=None, border=BOX, sz=11):
    c = ws[cell]
    c.value = val
    c.font = Font(bold=bold, size=sz)
    if fmt:   c.number_format = fmt
    if fill:  c.fill = fill
    c.alignment = align or (LFT if isinstance(val, str) else RGT)
    if border: c.border = border
    return c

# ---------- 1. limpiar ----------
for mr in list(ws.merged_cells.ranges):
    ws.unmerge_cells(str(mr))
for row in ws.iter_rows(min_row=1, max_row=320, min_col=1, max_col=40):
    for c in row:
        c.value = None
        c.fill = PatternFill()
        c.border = Border()
        c.font = Font()
        c.number_format = "General"
        c.alignment = Alignment()
for r in list(ws.row_dimensions.keys()):
    ws.row_dimensions[r].height = None
    ws.row_dimensions[r].hidden = False

# ---------- 2. anchos ----------
ANCHOS = {"A": 26, "B": 30, "C": 13, "D": 13, "E": 14, "F": 13, "G": 13,
          "H": 13, "I": 28, "J": 13, "K": 13, "L": 14}
for i in range(13, 27):   # M..Z
    ANCHOS[gl(i)] = 13
for k, v in ANCHOS.items():
    ws.column_dimensions[k].width = v

# ---------- 3. titulo ----------
ws.merge_cells("A1:L1")
B("A1", "Anexo de capacidad operativa", fill=VERDE, bold=True, align=CTR,
  border=Border(bottom=med), sz=20)
ws.row_dimensions[1].height = 32

# ---------- 4. horas laborables ----------
ws.merge_cells("A3:E3"); ws.merge_cells("H3:L3")
B("A3", "Instalación y alta de los planes ShopMetrics", fill=VERDE2, bold=True, align=CTR, sz=13)
B("H3", "Soporte y mantenimiento mensual por plan",     fill=VERDE2, bold=True, align=CTR, sz=13)

ws.merge_cells("A6:C6")
B("A6", "Cálculo de horas laborables de un empleado por mes", fill=VERDE2, bold=True, align=LFT, sz=12)
HL = [("Semanas por año", 52, "semanas", True),
      ("Semanas de licencias varias al año", 2, "semanas", True),
      ("Días festivos", 14, "días", True),
      ("Días laborales totales", "=(B7-B8)*5-B9", "días", False),
      ("Jornada laboral", 8, "horas", True),
      ("Horas laborables anuales por empleado", "=B10*B11", "horas", False)]
for i, (lab, val, uni, entrada) in enumerate(HL):
    r = 7 + i
    B(f"A{r}", lab, fill=AZUL)
    B(f"B{r}", val, fmt=NUM0, fill=AMAR if entrada else AZUL, bold=not entrada)
    B(f"C{r}", uni, fill=AZUL)
B("A13", "Consideraremos en promedio:", fill=AZUL, bold=True)
B("B13", 1800, fmt=NUM0, fill=AMAR, bold=True)
B("C13", "Horas laborables anuales / empleado", fill=AZUL)
B("A14", None, fill=AZUL)
B("B14", 150, fmt=NUM0, fill=AMAR, bold=True)
B("C14", "Horas laborables mensuales / empleado\n[HsLaboralesXMes]", fill=AMAR, bold=True, align=LFT)
ws.row_dimensions[14].height = 32

# ---------- 5. bloques de tareas ----------
INST = [
 ("Alta e instalación Plan Básico", 18, [
   (1, "Contacto y coordinación con el comercio",            0.08, 1),
   (2, "Alta del comercio en la plataforma",                 0.08, 1),
   (3, "Conexión con el sistema de punto de venta",          0.17, 1),
   (4, "Verificación de la calidad de los datos",            0.12, 1),
   (5, "Capacitación remota del comerciante",                0.05, 1)]),
 ("Alta e instalación Plan Vidriera", 37, [
   (1, "Coordinación de la visita y preparación del kit",    0.25, 1),
   (2, "Traslado al local",                                  0.50, 1),
   (3, "Montaje del sensor de vidriera",                     0.75, 1),
   (4, "Instalación del contador de puerta",                 0.50, 1),
   (5, "Conexión con el sistema de punto de venta",          0.25, 1),
   (6, "Prueba del embudo y calibración",                    0.50, 1),
   (7, "Capacitación del comerciante y cierre",              0.25, 1)]),
 ("Alta e instalación Plan Cadena", 56, [
   (1, "Relevamiento previo de la cadena",                   1.00, 1),
   (2, "Alta de la cadena y de sus sucursales",              0.50, 1),
   (3, "Traslado entre locales",                             0.50, 5),
   (4, "Montaje del sensor de vidriera",                     0.75, 5),
   (5, "Instalación del contador de puerta",                 0.40, 5),
   (6, "Conexión del punto de venta por local",              0.20, 5),
   (7, "Consolidación y prueba comparativa entre sucursales",0.75, 1),
   (8, "Capacitación del responsable de la cadena",          0.50, 1)])]

SOP = [
 ("Soporte mensual Plan Básico", 18, [
   (1, "Monitoreo automático de la ingesta de datos",        0.05, 1),
   (2, "Atención de consultas por canal remoto",             0.07, 1),
   (3, "Revisión y envío del reporte mensual",               0.03, 1)]),
 ("Soporte mensual Plan Vidriera", 37, [
   (1, "Monitoreo automático de la ingesta de datos",        0.05, 1),
   (2, "Control del estado del sensor de vidriera",          0.07, 1),
   (3, "Atención de consultas por canal remoto",             0.08, 1),
   (4, "Revisión y envío del reporte mensual",               0.05, 1)]),
 ("Soporte mensual Plan Cadena", 56, [
   (1, "Monitoreo automático de la ingesta (5 locales)",     0.04, 5),
   (2, "Control del estado de los sensores",                 0.03, 5),
   (3, "Atención de consultas del responsable de cadena",    0.10, 1),
   (4, "Reporte comparativo entre sucursales",               0.15, 1)])]

def bloque(titulo, r0, tareas, c0, precio_ref, es_inst):
    """c0 = indice de la primera columna (1 = A, 8 = H)."""
    L = lambda o: gl(c0 + o)
    ws.merge_cells(f"{L(0)}{r0}:{L(4)}{r0}")
    B(f"{L(0)}{r0}", titulo, fill=VERDE, bold=True, align=CTR, sz=12)
    hdr = ["N° de tarea", "Descripción", "Base horaria\n[BH]",
           "Cantidad requerida\n[Cant]", "Horas por tarea\nFórmula = Cant * BH"]
    for i, h in enumerate(hdr):
        B(f"{L(i)}{r0+1}", h, fill=HDR, bold=True, align=CTR)
    ws.row_dimensions[r0+1].height = 42
    for i in range(12):
        r = r0 + 2 + i
        t = tareas[i] if i < len(tareas) else (i + 1, None, None, None)
        B(f"{L(0)}{r}", t[0], fmt=NUM0, fill=AZUL, align=CTR)
        B(f"{L(1)}{r}", t[1], fill=AZUL)
        B(f"{L(2)}{r}", t[2], fmt=NUM2, fill=AMAR)
        B(f"{L(3)}{r}", t[3], fmt=NUM0, fill=AMAR)
        B(f"{L(4)}{r}", f"={L(3)}{r}*{L(2)}{r}", fmt=NUM2, fill=AZUL)
        ws.row_dimensions[r].height = 15
    rt = r0 + 14              # fila de totales (32 / 51 / 70)
    B(f"{L(3)}{rt}", "Horas totales por servicio\n[HsTotalesXServicio]", fill=AMAR, bold=True, align=LFT)
    B(f"{L(4)}{rt}", f"=SUM({L(4)}{r0+2}:{L(4)}{r0+13})", fmt=NUM2, fill=TOTFIL, bold=True)
    B(f"{L(3)}{rt+1}", "Capacidad operativa mensual por empleado\nFórmula = HsTotalesXServicio / HsLaboralesXMes",
      fill=AMAR, bold=True, align=LFT)
    B(f"{L(4)}{rt+1}", f"={L(4)}{rt}/$B$14", fmt="0.0000", fill=TOTFIL, bold=True)
    ws.row_dimensions[rt].height   = 32
    ws.row_dimensions[rt+1].height = 34
    if es_inst:   # costo de mano de obra de la instalación
        B(f"{L(0)}{rt}", "Costo de mano de obra", fill=AMAR, bold=True)
        B(f"{L(1)}{rt}", f"={L(4)}{rt}*$C$16", fmt=USD, fill=TOTFIL, bold=True)
        B(f"{L(0)}{rt+1}", "Precio de lista del alta", fill=AMAR, bold=True)
        B(f"{L(1)}{rt+1}", f"=Hipótesis!$C${precio_ref}", fmt=USD, fill=AZUL, bold=True)

for (ti, r0, ta), pref in zip(INST, [58, 59, 60]):
    bloque(ti, r0, ta, 1, pref, True)
for (ti, r0, ta), pref in zip(SOP, [61, 62, 63]):
    bloque(ti, r0, ta, 8, pref, False)

# costo horario del tecnico (usado por los bloques de instalacion)
ws.merge_cells("A16:B16")
B("A16", "Costo horario del técnico de instalación y soporte (USD)", fill=AMAR, bold=True, align=LFT)
B("C16", 6, fmt=USD, fill=AMAR, bold=True)
ws.row_dimensions[16].height = 18

# ---------- 6. analisis anual en horas hombre ----------
ws.merge_cells("A77:L77")
B("A77", "ANÁLISIS DE CAPACIDAD OPERATIVA ANUAL POR SERVICIO EN HORAS HOMBRE",
  fill=VERDE, bold=True, align=CTR, sz=13)
ws.row_dimensions[77].height = 24

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
         "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
PVCOL = ["E", "G", "I", "K", "M", "O", "Q", "S", "U", "W", "Y", "AA"]  # meses en Proy. ventas
CAP   = ["$E$32", "$E$51", "$E$70", "$L$32", "$L$51", "$L$70"]          # hs por unidad

def anio(r0, etiqueta, filas_altas, filas_abonos):
    B(f"A{r0}", etiqueta, fill=VERDE2, bold=True, align=CTR, sz=12)
    B(f"B{r0}", "Total anual", fill=VERDE2, bold=True, align=CTR)
    for i, m in enumerate(MESES):
        c = 3 + i * 2
        ws.merge_cells(start_row=r0, start_column=c, end_row=r0, end_column=c + 1)
        B(f"{gl(c)}{r0}", m, fill=VERDE2, bold=True, align=CTR)
    ws.row_dimensions[r0].height = 20
    for i in range(12):
        c = 3 + i * 2
        B(f"{gl(c)}{r0+1}",   "Unidades\nrequeridas",     fill=HDR, bold=True, align=CTR)
        B(f"{gl(c+1)}{r0+1}", "Horas Hombre\nrequeridas", fill=HDR, bold=True, align=CTR)
    B(f"A{r0+1}", "Servicio", fill=HDR, bold=True, align=CTR)
    B(f"B{r0+1}", "Horas Hombre", fill=HDR, bold=True, align=CTR)
    ws.row_dimensions[r0+1].height = 30

    pv = filas_altas + filas_abonos                       # 6 filas de Proy. ventas
    for j in range(6):
        r = r0 + 2 + j
        B(f"A{r}", f"=Hipótesis!B${58+j}", fill=AZUL)
        B(f"B{r}", "=" + "+".join(f"{gl(3+i*2+1)}{r}" for i in range(12)),
          fmt=NUM2, fill=TOTFIL, bold=True)
        for i in range(12):
            c = 3 + i * 2
            B(f"{gl(c)}{r}",   f"='Proy. ventas'!{PVCOL[i]}{pv[j]}", fmt=NUM0, fill=AZUL, align=CTR)
            B(f"{gl(c+1)}{r}", f"={gl(c)}{r}*{CAP[j]}",              fmt=NUM2, fill=AZUL)
        ws.row_dimensions[r].height = 15

    rt = r0 + 8
    B(f"A{rt}", "TOTALES", fill=TOTFIL, bold=True, border=BOXB)
    B(f"B{rt}", "=" + "+".join(f"{gl(3+i*2+1)}{rt}" for i in range(12)),
      fmt=NUM2, fill=TOTFIL, bold=True, border=BOXB)
    for i in range(12):
        c = 3 + i * 2
        B(f"{gl(c)}{rt}",   f"=SUM({gl(c)}{r0+2}:{gl(c)}{r0+7})",     fmt=NUM0, fill=TOTFIL, bold=True, align=CTR, border=BOXB)
        B(f"{gl(c+1)}{rt}", f"=SUM({gl(c+1)}{r0+2}:{gl(c+1)}{r0+7})", fmt=NUM2, fill=TOTFIL, bold=True, border=BOXB)

    rn = rt + 1
    B(f"A{rn}", "Técnicos necesarios", fill=AMAR, bold=True)
    B(f"B{rn}", f"=MAX({gl(4)}{rn}:{gl(26)}{rn})", fmt=NUM0, fill=AMAR, bold=True, align=CTR)
    for i in range(12):
        c = 3 + i * 2
        ws.merge_cells(start_row=rn, start_column=c, end_row=rn, end_column=c + 1)
        B(f"{gl(c)}{rn}", f"=ROUNDUP({gl(c+1)}{rt}/$B$14,0)", fmt=NUM0, fill=AMAR, bold=True, align=CTR)
    ws.row_dimensions[rn].height = 18

anio(78,  "CAPACIDAD OPERATIVA AÑO 2026", [19, 20, 21],    [22, 23, 24])
anio(89,  "CAPACIDAD OPERATIVA AÑO 2027", [82, 83, 84],    [85, 86, 87])
anio(100, "CAPACIDAD OPERATIVA AÑO 2028", [144, 145, 146], [147, 148, 149])

# ---------- 7. conclusion ----------
ws.merge_cells("A112:L112")
B("A112", "CONCLUSIÓN", fill=VERDE, bold=True, align=CTR, sz=13)
ws.row_dimensions[112].height = 22

CH = ["Año", "Horas Hombre\ntotales del año", "Pico mensual de\nHoras Hombre",
      "Horas laborables\npor empleado / mes", "Técnicos\nnecesarios",
      "Costo de mano de obra\nde instalación y soporte"]
for i, h in enumerate(CH):
    B(f"{gl(1+i)}113", h, fill=HDR, bold=True, align=CTR)
ws.row_dimensions[113].height = 42
for i, (a, rt, rn) in enumerate([(2026, 86, 87), (2027, 97, 98), (2028, 108, 109)]):
    r = 114 + i
    B(f"A{r}", a, fmt=NUM0, fill=AZUL, align=CTR)
    B(f"B{r}", f"=B{rt}", fmt=NUM2, fill=AZUL)
    B(f"C{r}", f"=MAX(D{rt},F{rt},H{rt},J{rt},L{rt},N{rt},P{rt},R{rt},T{rt},V{rt},X{rt},Z{rt})",
      fmt=NUM2, fill=AZUL)
    B(f"D{r}", "=$B$14", fmt=NUM0, fill=AZUL, align=CTR)
    B(f"E{r}", f"=B{rn}", fmt=NUM0, fill=TOTFIL, bold=True, align=CTR)
    B(f"F{r}", f"=B{r}*$C$16", fmt=USD, fill=AZUL)

ws.merge_cells("A118:L122")
B("A118",
  "La capacidad operativa se dimensiona desde la demanda. Cada plan tiene un consumo horario propio, obtenido de la suma de las "
  "tareas que requiere su alta, y un consumo mensual de soporte que se sostiene mientras el comercio permanece abonado. Multiplicando "
  "esas bases horarias por las unidades de la proyección de ventas se obtienen las horas hombre requeridas mes a mes.\n\n"
  "El cociente entre las horas hombre del mes de mayor exigencia y las horas laborables mensuales de un empleado determina la "
  "dotación técnica necesaria en cada año. El mes en que ese cociente supera la unidad es el mes en que aparece la incorporación "
  "en la estructura de costos de recursos humanos: la contratación no se decide por criterio, se desprende del volumen proyectado.\n\n"
  "El costo de mano de obra de instalación y soporte que surge de este anexo permite además contrastar el precio de lista de cada "
  "alta contra el costo real de ejecutarla, y verificar que la tarifa cubre el esfuerzo técnico comprometido.",
  fill=GRIS, align=Alignment(horizontal="justify", vertical="top", wrap_text=True), border=BOX)
for r in range(118, 123):
    ws.row_dimensions[r].height = 30

ws.sheet_view.selection[0].activeCell = "A1"
ws.sheet_view.selection[0].sqref = "A1"
ws.sheet_view.showGridLines = False
wb.save(P)
print("OK - Anexo reconstruido")
