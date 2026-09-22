# -*- coding: utf-8 -*-
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
P = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wb = openpyxl.load_workbook(P); ws = wb["Anexo capacidad operativa"]

# titulo del analisis a lo ancho de toda la tabla (A:Z)
for m in [str(x) for x in ws.merged_cells.ranges]:
    if m.startswith("A77:"):
        ws.unmerge_cells(m)
ws.merge_cells("A77:Z77")

# alturas de los encabezados de cada año
for r0 in (78, 89, 100):
    ws.row_dimensions[r0].height     = 30   # nombre del año + meses
    ws.row_dimensions[r0 + 1].height = 46   # "Unidades / Horas Hombre requeridas"
    ws.row_dimensions[r0 + 8].height = 18   # TOTALES
    ws.row_dimensions[r0 + 9].height = 18   # Tecnicos necesarios

# etiqueta larga de capacidad operativa: un poco mas de alto
for rt in (33, 52, 71):
    ws.row_dimensions[rt].height = 58

# conclusion: año sin separador de miles, encabezado mas alto
ws.row_dimensions[113].height = 52
for r in (114, 115, 116):
    ws[f"A{r}"].number_format = "0"

wb.save(P); print("OK fix4")
