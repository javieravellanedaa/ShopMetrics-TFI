# -*- coding: utf-8 -*-
import openpyxl, sys, io
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

P = r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wb = openpyxl.load_workbook(P); ws = wb["Anexo capacidad operativa"]

AMAR = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)
LFT  = Alignment(horizontal="left", vertical="center", wrap_text=True)

TXT = ["Horas totales por servicio\n[HsTotalesXServicio]",
       "Capacidad operativa mensual por empleado\nFórmula = HsTotalesXServicio / HsLaboralesXMes"]
USADAS = {18: 5, 37: 7, 56: 8}

# alturas del bloque de horas laborables
ws.row_dimensions[13].height = 46
ws.row_dimensions[14].height = 52

for r0, n in USADAS.items():
    for i in range(12):
        ws.row_dimensions[r0 + 2 + i].height = 30 if i < n else 16
    rt = r0 + 14
    for k, rr in enumerate((rt, rt + 1)):
        for a, b in (("C", "D"), ("J", "K")):
            rng = f"{a}{rr}:{b}{rr}"
            ws[f"{b}{rr}"].value = None          # limpiar ANTES de fusionar
            if rng not in [str(m) for m in ws.merged_cells.ranges]:
                ws.merge_cells(rng)
            c = ws[f"{a}{rr}"]
            c.value = TXT[k]; c.font = Font(bold=True)
            c.fill = AMAR; c.alignment = LFT; c.border = BOX
    ws.row_dimensions[rt].height     = 40
    ws.row_dimensions[rt + 1].height = 52

wb.save(P)

# ---- verificacion ----
wb2 = openpyxl.load_workbook(P); w2 = wb2["Anexo capacidad operativa"]
mg = [str(m) for m in w2.merged_cells.ranges]
for r0 in (18, 37, 56):
    rt = r0 + 14
    print(f"bloque {r0}: merges {[x for x in mg if x.startswith(('C%d'%rt,'D%d'%rt,'J%d'%rt,'K%d'%rt))]}")
    print(f"   C{rt}={repr(w2[f'C{rt}'].value)[:34]}  D{rt}={repr(w2[f'D{rt}'].value)}  alt={w2.row_dimensions[rt].height}")
    print(f"   vacias h={[w2.row_dimensions[r0+2+i].height for i in range(12)]}")
