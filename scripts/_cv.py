# -*- coding: utf-8 -*-
import openpyxl,sys,io
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding="utf-8",errors="replace")
p=r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wf=openpyxl.load_workbook(p,data_only=False); wv=openpyxl.load_workbook(p,data_only=True)
for sn in ["Costos variables","Costos RRHH"]:
    f,v=wf[sn],wv[sn]; print("="*60); print(sn, f.max_row, f.max_column)
    for r in range(1, min(f.max_row,22)+1):
        fila=[]
        for c in range(1, min(f.max_column,11)+1):
            cf=f.cell(r,c); cv=v.cell(r,c)
            if cf.value is None and cv.value is None: continue
            s=f"{cf.coordinate}="+(f"{cf.value[:40]}->{cv.value}" if isinstance(cf.value,str) and cf.value.startswith("=") else repr(cf.value)[:48])
            fila.append(s)
        if fila: print(f"[{r:>2}]"," | ".join(fila))
