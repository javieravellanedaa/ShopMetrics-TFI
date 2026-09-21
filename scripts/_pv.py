# -*- coding: utf-8 -*-
import openpyxl,sys,io
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding="utf-8",errors="replace")
p=r"D:\Javier\UAI\UAI - SAP\STF\STF\Primera entrega\Presupuesto financiero ShopMetrics V1.xlsx"
wf=openpyxl.load_workbook(p,data_only=False); wv=openpyxl.load_workbook(p,data_only=True)
f,v=wf["Proy. ventas"],wv["Proy. ventas"]
from openpyxl.utils import get_column_letter as gl
for r in list(range(14,36)):
    fila=[]
    for c in range(1, 29):
        cf=f.cell(r,c); cv=v.cell(r,c)
        if cf.value is None and cv.value is None: continue
        s=f"{cf.coordinate}="
        if isinstance(cf.value,str) and cf.value.startswith("="): s+=f"{cv.value}"
        else: s+=repr(cf.value)[:55]
        fila.append(s)
    if fila: print(f"[{r:>3}]"," | ".join(fila))
